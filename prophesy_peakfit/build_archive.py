"""Package PROPHESY simulations with known peak parameters for ML use.

Only the simulated data.xlsx files with matching meta_hν_* sheets are used.
Every Snoisy repeat scan inherits the generating parameters of its condition.
Experimental fits, older unlabeled simulations, and IGOR reprocessings are
intentionally excluded from this supervised dataset.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
from collections import defaultdict
from pathlib import Path
from zipfile import ZIP_STORED, ZipFile

import numpy as np
from openpyxl import load_workbook


GRIDS = {
    "C1s": np.linspace(271.0, 291.0, 401, dtype=np.float64),
    "S2p": np.linspace(165.0, 174.5, 191, dtype=np.float64),
    # The water workbooks cover either 526–542 or 528–542 eV. Use their
    # common measured interval, avoiding fabricated values outside a scan.
    "O1s": np.linspace(528.0, 542.0, 281, dtype=np.float64),
}
EXPECTED_BASE_COUNTS = {"C1s": 1020, "S2p": 100, "O1s": 380}
EXPECTED_SWEEP_COUNTS = {"C1s": 46000, "S2p": 5000, "O1s": 0}


def table(sheet, wanted: set[str] | None = None) -> dict[str, np.ndarray]:
    rows = sheet.values
    header = next(rows)
    names = [str(value) if value is not None else "" for value in header]
    positions = [(i, name) for i, name in enumerate(names) if name and (wanted is None or name in wanted)]
    columns: dict[str, list[float]] = {name: [] for _, name in positions}
    for row in rows:
        for i, name in positions:
            value = row[i] if i < len(row) else None
            if value is None:
                raise ValueError(f"Missing {name} value in {sheet.title}")
            columns[name].append(float(value))
    return {name: np.asarray(values, dtype=np.float64) for name, values in columns.items()}


def classify(path: Path, meta: dict[str, np.ndarray]) -> str:
    if "peak_mode_liq" in meta:
        return "O1s"
    return "S2p" if "S2p" in path.parts else "C1s"


def sampled(values: np.ndarray, source_axis: np.ndarray, target_axis: np.ndarray) -> np.ndarray:
    order = np.argsort(source_axis)
    ordered_axis = source_axis[order]
    if not np.all(np.diff(ordered_axis) > 0):
        raise ValueError("Binding-energy axis contains duplicate values")
    if target_axis[0] < ordered_axis[0] - 1e-5 or target_axis[-1] > ordered_axis[-1] + 1e-5:
        raise ValueError("Target grid extends beyond a measured spectrum")
    return np.interp(target_axis, ordered_axis, values[order]).astype(np.float32)


def records(path: Path, data_root: Path, sheet_name: str, meta_name: str, book) -> list[dict]:
    sweep_names = sorted(
        (name for name in next(book[sheet_name].values)
         if isinstance(name, str) and name.startswith("Snoisy_") and name[7:].isdigit()),
        key=lambda name: int(name[7:]),
    )
    wanted = {"Be", "Snoisy", "Sbg", "SpectrumA_1", "SpectrumA_1_gas", *sweep_names}
    spectrum = table(book[sheet_name], wanted)
    meta = table(book[meta_name])
    region = classify(path, meta)
    axis = GRIDS[region]
    source_axis = spectrum["Be"]
    noisy = sampled(spectrum["Snoisy"], source_axis, axis)
    background = sampled(spectrum["Sbg"], source_axis, axis)
    liquid_or_target = sampled(spectrum["SpectrumA_1"], source_axis, axis)

    if region == "O1s":
        gas = sampled(spectrum["SpectrumA_1_gas"], source_axis, axis)
        clean = liquid_or_target + gas
        centers = np.asarray([meta["peak_mode_liq"][0], meta["peak_mode_gas"][0]])
        sigmas = np.asarray([meta["peak_width_liq"][0], meta["peak_width_gas"][0]])
        areas = np.asarray([np.trapezoid(liquid_or_target, axis), np.trapezoid(gas, axis)])
        fractions = areas / areas.sum()
    else:
        clean = liquid_or_target
        centers = meta["peak_mode"]
        sigmas = meta["peak_width"]
        fractions = meta["peak_probability"]
        areas = fractions * np.trapezoid(clean, axis)

    if not (len(centers) == len(sigmas) == len(fractions) == len(areas)):
        raise ValueError(f"Peak label lengths differ in {path}: {meta_name}")
    if not np.isclose(fractions.sum(), 1.0, atol=1e-4):
        raise ValueError(f"Peak fractions do not sum to one in {path}: {meta_name}")
    if np.any(sigmas <= 0) or np.any(centers < axis[0]) or np.any(centers > axis[-1]):
        raise ValueError(f"Peak center or width outside the usable grid in {path}: {meta_name}")
    if not all(np.isfinite(a).all() for a in (noisy, background, clean, centers, sigmas, areas)):
        raise ValueError(f"Non-finite spectrum or label in {path}: {sheet_name}")

    relative = path.relative_to(data_root).as_posix()
    base = {
        "region": region,
        "sample_id": f"{relative}#{sheet_name}",
        "condition_id": f"{relative}#{sheet_name}",
        "scan": "Snoisy",
        "source_workbook": relative,
        "photon_energy_eV": float(meta["hν"][0]),
        "noisy_counts": noisy,
        "clean_counts": clean.astype(np.float32),
        "background_counts": background,
        "peak_center_eV": centers.astype(np.float32),
        "peak_sigma_eV": sigmas.astype(np.float32),
        "peak_fraction": fractions.astype(np.float32),
        "peak_area_counts_eV": areas.astype(np.float32),
    }
    result = [base]
    for name in sweep_names:
        result.append({**base, "sample_id": f"{relative}#{sheet_name}/{name}",
                       "scan": name, "noisy_counts": sampled(spectrum[name], source_axis, axis)})
    return result


def split_by_workbook(records: list[dict], seed: int) -> None:
    groups = sorted({item["source_workbook"] for item in records})
    rng = np.random.default_rng(seed)
    shuffled = [groups[i] for i in rng.permutation(len(groups))]
    n_train = int(0.70 * len(groups))
    n_validation = int(0.15 * len(groups))
    labels = (["train"] * n_train + ["validation"] * n_validation
              + ["test"] * (len(groups) - n_train - n_validation))
    group_split = dict(zip(shuffled, labels))
    for item in records:
        item["split"] = group_split[item["source_workbook"]]


def make_archive(data_root: Path, output: Path, seed: int = 42) -> dict:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for path in sorted(data_root.rglob("data.xlsx")):
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            names = set(book.sheetnames)
            for sheet_name in sorted(name for name in names if name.startswith("hν_")):
                meta_name = "meta_" + sheet_name
                if meta_name in names:
                    for item in records(path, data_root, sheet_name, meta_name, book):
                        grouped[item["region"]].append(item)
        finally:
            book.close()

    base_counts = {region: sum(item["scan"] == "Snoisy" for item in grouped[region]) for region in GRIDS}
    sweep_counts = {region: len(grouped[region]) - base_counts[region] for region in GRIDS}
    if base_counts != EXPECTED_BASE_COUNTS or sweep_counts != EXPECTED_SWEEP_COUNTS:
        raise ValueError(f"Expected {EXPECTED_BASE_COUNTS} base and {EXPECTED_SWEEP_COUNTS} "
                         f"repeat spectra, found {base_counts} and {sweep_counts}")
    for i, region in enumerate(GRIDS):
        split_by_workbook(grouped[region], seed + i)

    output.parent.mkdir(parents=True, exist_ok=True)
    manifest_rows = []
    summary = {"schema": "prophesy_peakfit_v1", "seed": seed, "regions": {}}
    with ZipFile(output, "w", compression=ZIP_STORED) as archive:
        for region, axis in GRIDS.items():
            items = sorted(grouped[region], key=lambda item: item["sample_id"])
            fields = (
                "noisy_counts", "clean_counts", "background_counts", "peak_center_eV",
                "peak_sigma_eV", "peak_fraction", "peak_area_counts_eV",
            )
            arrays = {key: np.stack([item[key] for item in items]) for key in fields}
            arrays["axis_eV"] = axis.astype(np.float32)
            for key in ("sample_id", "condition_id", "scan", "source_workbook", "split"):
                arrays[key] = np.asarray([item[key] for item in items], dtype=np.str_)
            arrays["photon_energy_eV"] = np.asarray(
                [item["photon_energy_eV"] for item in items], dtype=np.float32
            )
            binary = io.BytesIO()
            np.savez_compressed(binary, **arrays)
            archive.writestr(f"{region}.npz", binary.getvalue())
            splits, split_counts = np.unique(arrays["split"], return_counts=True)
            summary["regions"][region] = {
                "spectra": len(items), "base_spectra": base_counts[region],
                "repeat_scans": sweep_counts[region], "energy_points": len(axis),
                "peaks_per_spectrum": arrays["peak_center_eV"].shape[1],
                "splits": dict(zip(splits.tolist(), split_counts.tolist())),
            }
            for item in items:
                manifest_rows.append({
                    "sample_id": item["sample_id"], "condition_id": item["condition_id"],
                    "scan": item["scan"], "region": region,
                    "source_workbook": item["source_workbook"],
                    "photon_energy_eV": item["photon_energy_eV"], "split": item["split"],
                })

        csv_buffer = io.StringIO()
        writer = csv.DictWriter(csv_buffer, fieldnames=list(manifest_rows[0]))
        writer.writeheader()
        writer.writerows(manifest_rows)
        archive.writestr("manifest.csv", csv_buffer.getvalue())
        archive.writestr("metadata.json", json.dumps(summary, indent=2) + "\n")
        archive.writestr("README.txt", (
            "PROPHESY peak-fitting dataset, schema prophesy_peakfit_v1.\n"
            "Each NPZ holds one spectral region and a common binding-energy grid.\n"
            "noisy_counts is the measured-like input; clean_counts is the noise-free\n"
            "peak signal without background. peak_sigma_eV is Gaussian sigma, not FWHM.\n"
            "C1s/S2p peak_fraction is the generator's cross-section mixture weight;\n"
            "peak_area_counts_eV follows the repository's fraction × total-area rule.\n"
            "O1s fractions and areas come from the separately saved liquid and gas\n"
            "signals. Its common measured grid is 528–542 eV. The ZIP contains\n"
            "primary and repeated noisy scans; repeats share condition labels.\n"
            "Keep all scans from a workbook in one data split. Experimental fitting\n"
            "tables and older unlabeled simulations are excluded.\n"
        ))
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=Path(__file__).resolve().parents[2] / "PROPHESY" / "data")
    parser.add_argument("--output", type=Path, default=Path("datasets/prophesy_peakfit.zip"))
    parser.add_argument("--seed", type=int, default=42)
    args = parser.parse_args()
    result = make_archive(args.source.resolve(), args.output.resolve(), args.seed)
    print(json.dumps({"archive": str(args.output.resolve()), **result}, indent=2))


if __name__ == "__main__":
    main()
